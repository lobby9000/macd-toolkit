"""
Cross-check the Excel workbook against macd_toolkit.py.

The workbook in excel/plantilla_MACD.xlsx computes the same indicators with
linked spreadsheet formulas. This script recalculates it with LibreOffice in
headless mode, runs the Python implementation on the same closing prices and
reports the largest absolute difference for every quantity, plus whether both
sides flag the same MACD/signal crossovers.

    python scripts/excel_crosscheck.py
    python scripts/excel_crosscheck.py --tolerance 1e-9 --soffice /usr/bin/soffice

Exit status is 0 when every difference is within the tolerance and the
crossovers match, and 1 otherwise.
"""

from __future__ import annotations

import argparse
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import macd_toolkit as mt  # noqa: E402

FIRST_ROW, LAST_ROW = 9, 112          # calculation rows in every sheet (t = 1..104)


def count_formulas(path: Path) -> int:
    wb = openpyxl.load_workbook(path)
    return sum(
        1
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    )


def recalculate(path: Path, soffice: str) -> Path:
    """Open and re-save the workbook with LibreOffice so every formula has a value."""
    out_dir = Path(tempfile.mkdtemp(prefix="macd-xlsx-"))
    src = out_dir / "in" / path.name
    src.parent.mkdir()
    shutil.copy(path, src)
    subprocess.run(
        [soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out_dir), str(src)],
        check=True, capture_output=True, timeout=300,
    )
    return out_dir / path.name


def column(ws, letter: str) -> pd.Series:
    values = [ws[f"{letter}{r}"].value for r in range(FIRST_ROW, LAST_ROW + 1)]
    return pd.Series([v if isinstance(v, (int, float)) else np.nan for v in values], dtype="float64")


def max_diff(excel: pd.Series, python: pd.Series) -> tuple[int, float]:
    both = excel.notna() & python.notna()
    if (excel.notna() != python.notna()).any():
        return int(both.sum()), float("inf")      # warm-up periods disagree
    return int(both.sum()), float((excel[both] - python[both]).abs().max())


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description=__doc__.splitlines()[1])
    p.add_argument("--workbook", type=Path, default=ROOT / "excel" / "plantilla_MACD.xlsx")
    p.add_argument("--soffice", default=shutil.which("soffice") or shutil.which("libreoffice"))
    p.add_argument("--tolerance", type=float, default=1e-9)
    args = p.parse_args(argv)
    if not args.soffice:
        print("LibreOffice (soffice) is needed to recalculate the workbook.", file=sys.stderr)
        return 2

    n_formulas = count_formulas(args.workbook)
    wb = openpyxl.load_workbook(recalculate(args.workbook, args.soffice), data_only=True)

    close = column(wb["MACD_EMA"], "C")
    ema = mt.macd(close)
    sma = mt.macd(close, mt.MACDConfig(kind="sma"))
    checks = {
        "EMA 12": (column(wb["MACD_EMA"], "D"), ema["ma_fast"]),
        "EMA 26": (column(wb["MACD_EMA"], "E"), ema["ma_slow"]),
        "MACD": (column(wb["MACD_EMA"], "F"), ema["macd"]),
        "Signal": (column(wb["MACD_EMA"], "G"), ema["signal"]),
        "Histogram": (column(wb["MACD_EMA"], "H"), ema["hist"]),
        "PPO (%)": (100 * column(wb["MACD_EMA"], "L"), mt.ppo(close)),
        "MACD, SMA variant": (column(wb["MACD_SMA"], "F"), sma["macd"]),
        "Signal, SMA variant": (column(wb["MACD_SMA"], "G"), sma["signal"]),
        "RSI 14, class method": (column(wb["RSI"], "J"), mt.rsi(close, 14, "simple")),
        "RSI 14, Wilder": (column(wb["RSI"], "M"), mt.rsi(close, 14, "wilder")),
    }

    # Worked example: MACD(3, 6, 3) on 12 prices, small enough to check by hand.
    hand = wb["Ejemplo_a_mano"]
    prices = pd.Series([hand[f"B{r}"].value for r in range(9, 21)], dtype="float64")
    small = mt.macd(prices, mt.MACDConfig(fast=3, slow=6, signal=3))
    hand_col = lambda c: pd.Series(  # noqa: E731
        [hand[f"{c}{r}"].value if isinstance(hand[f"{c}{r}"].value, (int, float)) else np.nan
         for r in range(9, 21)], dtype="float64")
    checks["Hand example MACD(3, 6, 3)"] = (hand_col("E"), small["macd"])
    checks["Hand example signal"] = (hand_col("F"), small["signal"])

    print(f"Workbook: {args.workbook.relative_to(ROOT)} ({n_formulas:,} formulas)\n")
    print(f"{'Quantity':<28}{'Values':>8}{'Max |Excel - Python|':>24}")
    worst = 0.0
    for name, (xl, py) in checks.items():
        n, d = max_diff(xl, py.reset_index(drop=True))
        worst = max(worst, d)
        print(f"{name:<28}{n:>8}{d:>24.3e}")

    labels = {"Cruce alcista": "cruce_alcista", "Cruce bajista": "cruce_bajista"}
    excel_events = {
        (r - FIRST_ROW, labels[v])
        for r in range(FIRST_ROW, LAST_ROW + 1)
        if (v := wb["MACD_EMA"][f"I{r}"].value) in labels
    }
    events = mt.crossovers(ema.reset_index(drop=True))
    python_events = {
        (int(row.fecha), row.evento)
        for row in events.itertuples()
        if row.evento in ("cruce_alcista", "cruce_bajista")
    }
    same = excel_events == python_events
    print(f"\nMACD/signal crossovers: Excel {len(excel_events)}, Python {len(python_events)}, "
          f"{'identical' if same else 'DIFFERENT'}")
    ok = same and worst <= args.tolerance
    print(f"Largest difference: {worst:.1e} (tolerance {args.tolerance:.0e}) -> {'PASS' if ok else 'FAIL'}")
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
